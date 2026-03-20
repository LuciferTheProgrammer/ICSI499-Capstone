from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches
from docx import Document
import pymupdf

# For Appendix. To insert screenshot of image after section header "SCREENSHOT OF OPEN PORTS TABLE"
def add_new_paragraph(paragraph):
    inserted = OxmlElement("w:p")
    paragraph._p.addnext(inserted)
    return Paragraph(inserted, paragraph._parent)

# To populate the Appendix section of the Findings Report.
def appendix(technical_report: str, findings_report: str) -> None:
    title = "Appendix B: Host Discovery (Opened Ports)"
    image_container = "../Reports/OPEN_PORTS.PNG"
    pdf_p = pymupdf.open(technical_report)
    target = None
    table_container = None
    len_pdf_p = len(pdf_p)
    for i in range(len_pdf_p):
        page = pdf_p[i]
        text_data = page.get_text()
        if title in text_data:
            target = page
            target_table = page.find_tables()
            if target_table and target_table.tables:
                max_table = max(target_table.tables, key = lambda x: (x.bbox[2] - x.bbox[0]) * (x.bbox[3] - x.bbox[1]))
                x0, y0, x1, y1 = max_table.bbox
                table_container = pymupdf.Rect(x0 - 3, y0 - 3, x1 + 3, y1 + 3)
            else:
                table_header = page.search_for("IP Address")
                protocol = page.search_for("Protocol")
                if protocol and table_header:
                    top = table_header[0].y0 - 3
                    bottom = table_header[0].y1 + 170
                    left_point = table_header[0].x0 - 3
                    right_point = page.rect.width - 36
                    table_container = pymupdf.Rect(left_point, top, right_point, bottom)
                else:
                    table_container = pymupdf.Rect(36, 150, page.rect.width - 36, 320)
            break
    if target is None:
        print(f"The corresponding section: {title} was not found in {technical_report}")
        pdf_p.close()
        return
    image_zoomed = pymupdf.Matrix(2.0, 2.0)
    pixels = target.get_pixmap(matrix=image_zoomed, clip=table_container ,alpha=False)
    pixels.save(image_container)
    pdf_p.close()
    print(f"✅ Screenshot of Open Ports Table was a success.")
    doc = Document(findings_report)
    idx_screenshot = None
    position = None
    text_1 = "OPEN PORTS | EXTERNAL NETWORK TESTING"
    text_2 = "SCREENSHOT OF OPEN PORTS TABLE"
    for i, p in enumerate(doc.paragraphs):
        cur_data = p.text.upper()
        if text_1 in cur_data:
            position = i
        if text_2 in cur_data:
            idx_screenshot = i
            break
    if position is None:
        print(f"❌ {text_1} can't be found")
        return
    if idx_screenshot is None:
        print(f"❌ {text_2} can't be found")
        return
    data_container = doc.paragraphs[idx_screenshot]
    image_struc = add_new_paragraph(data_container)
    image = image_struc.add_run()
    image.add_picture(image_container, width=Inches(6.5))
    image_struc.paragraph_format.space_before = Pt(0)
    image_struc.paragraph_format.space_after = Pt(6)
    print(f"✅ Open Ports Table was pasted over Findings Report successfully.")
    doc.save(findings_report)
    print("✅ Findings Report was successfully saved.")