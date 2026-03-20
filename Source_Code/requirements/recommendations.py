from docx.shared import Pt, Inches
import pandas as pd
import pdfplumber
from docx import Document
import re

def process_titles_rec(x ,y):
    x = y.lower()
    x = x.replace("\n", " ")
    x = " ".join(x.split())
    x = x.replace("–", "-").replace("—", "-")
    x = x.strip()
    return x

# Populates the Recommendations section in the Findings Report.
def recommendations(recommendation_csv: str, technical_report: str, findings_report: str, environment: str) -> None:
    #environment = input("Please select the environment type: 'Internal' or 'External': ")
    data_container = pd.read_excel(recommendation_csv, sheet_name=environment)
    recommendation_map = {}
    for _, i in data_container.iterrows():
        findings_title = i["Finding Title"]
        recommendation_data = i["Recommendation"]
        if findings_title is None:
            continue
        if recommendation_data is None:
            continue
        findings_title = str(findings_title).strip()
        recommendation_data = str(recommendation_data).strip()
        if findings_title == "":
            continue
        if recommendation_data == "" or recommendation_data == "nan":
            continue
        cleaned_titles = ""
        cleaned_titles = process_titles_rec(cleaned_titles, findings_title)
        recommendation_map[cleaned_titles] = recommendation_data
    pages_content = []
    technical_finds = []
    with pdfplumber.open(technical_report) as pdf:
        for page in pdf.pages:
            page_content = page.extract_text()
            if not page_content:
                continue
            pages_content.append(page_content)
            for entry in page_content.splitlines():
                normalized = entry.strip()
                if normalized == "":
                    continue
                matching = re.match(r"^(CRITICAL|HIGH|MEDIUM|LOW|INFORMATIONAL)\s+(.+)$", normalized, re.IGNORECASE)
                if matching:
                    status = matching.group(1).title()
                    finding_title = matching.group(2).strip()
                    if finding_title == "":
                        continue
                    cleaned_title = ""
                    cleaned_title = process_titles_rec(cleaned_title, finding_title)
                    technical_finds.append({"severity": status, "finding_title": finding_title, "cleaned_title": cleaned_title})
    recommendations = []
    viewed = set()
    status_counter = {"Critical": 0, "High": 0, "Medium": 0, "Low": 0, "Informational": 0}
    for finding in technical_finds:
        processed_title = finding["cleaned_title"]
        if processed_title in viewed:
            continue
        viewed.add(processed_title)
        if processed_title not in recommendation_map:
            #print(f"No recommendation found for {finding['finding_title']}")
            continue
        status_counter[finding["severity"]] = status_counter[finding["severity"]] + 1
        match = recommendation_map[processed_title]
        recommendations.append({"severity": finding["severity"], "tracker_finding": status_counter[finding["severity"]], "finding_title": finding["finding_title"], "recommendation_data": match})
    doc = Document(findings_report)
    position = None
    landing_section = "DISCRETIONARY REMEDIATION"
    for index, p in enumerate(doc.paragraphs):
        if landing_section in p.text.upper():
            position = index
            break
    if position is None:
        print(f"Recommendation section not found in {findings_report}")
        return
    len_paragraphs = len(doc.paragraphs)
    increment_position = position + 1
    if len_paragraphs > increment_position:
        new_position = doc.paragraphs[increment_position]
    else:
        new_position = doc.paragraphs[position]

    for i in recommendations:
        status_holder = i["severity"]
        tracker_holder = i["tracker_finding"]
        data_holder = i["recommendation_data"]
        left_component = data_holder
        right_component = ""
        if " - " in data_holder:
            left_component, right_component = data_holder.split(" - ", 1)
        elif " – " in data_holder:
            left_component, right_component = data_holder.split(" – ", 1)
        elif "-" in data_holder:
            left_component, right_component = data_holder.split("-", 1)
        left_component = left_component.strip()
        right_component = right_component.strip()
        body_cont1 = new_position.insert_paragraph_before()
        body_cont1.style = "List Number 2"
        body_cont2 = body_cont1.add_run(left_component)
        body_cont2.bold = True
        body_cont2.font.name = "Corbel"
        body_cont2.font.size = Pt(12)
        if right_component:
            normalized_def = body_cont1.add_run(" - " + right_component)
            normalized_def.font.name = "Corbel"
            normalized_def.font.size = Pt(12)
        rec_cont1 = new_position.insert_paragraph_before()
        rec_cont2 = rec_cont1.add_run(f"Refer to {status_holder} Finding {tracker_holder}")
        rec_cont2.font.name = "Corbel"
        rec_cont2.font.size = Pt(12)
        rec_cont2.underline = True
        rec_cont1.paragraph_format.left_indent = Inches(0.5)
        rec_cont1.paragraph_format.first_line_indent = Inches(0)
        rec_cont1.paragraph_format.space_before = Pt(0)
        rec_cont1.paragraph_format.space_after = Pt(6)
    doc.save(findings_report)
    print("✅ Recommendations was saved.")