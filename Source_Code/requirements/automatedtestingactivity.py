from camelot.io import read_pdf
from docx import Document
from docx.shared import Pt

"""
Find all testing from the activity report, and append it to the findings report.
"""
def automated_testing_activity(activity_report: str, findings_report: str, page_range: str) -> None:
    table_container = read_pdf(activity_report, pages=page_range, flavor="lattice")
    activity_log = []
    for entry in table_container:
        sample = entry.df
        for event in sample.values.tolist():
            string1 = event[0].replace("\n", "")
            string2 = event[1]
            string3 = event[2]
            cleaned = string1 + " " + string2 + " " + string3
            activity_log.append(cleaned)
    print("✅ All events copied.")
    doc_holder = Document(findings_report)
    section = "AUTOMATED TESTING ACTIVITY"
    for i, paragraph in enumerate(doc_holder.paragraphs):
        if section.lower() in paragraph.text.lower():
            for entry in reversed(activity_log):
                container = doc_holder.paragraphs[i + 1].insert_paragraph_before()
                container.style = "Activity Bullet"
                execute  = container.add_run(entry)
                execute.font.name = "Corbel"
                execute.font.size = Pt(8.5)
            break
    print("✅ All events pasted.")
    doc_holder.save(findings_report)
    print("✅ Document saved.")