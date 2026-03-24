from requirements.severitycounter import severity_counter
import pandas as pd
from docx import Document
from docx.shared import Pt

def normalize_title(title: str) -> str:
    return " ".join(str(title).strip().lower().split())

def load_findings_lookup(details_excel: str, sheet_name: str = "External") -> dict:
    df = pd.read_excel(details_excel, sheet_name=sheet_name)

    # clean column names
    df.columns = [str(col).strip() for col in df.columns]

    lookup = {}

    for _, row in df.iterrows():
        finding_title = str(row.get("Finding Title", "")).strip()
        if not finding_title:
            continue

        key = normalize_title(finding_title)

        lookup[key] = {
            "title": finding_title,
            "description": str(row.get("Finding Description", "")).strip(),
            "risk": str(row.get("Risk Level", "")).strip(),
            "recommendation": str(row.get("Recommendation", "")).strip(),
        }

    return lookup

def finding_details(technical_report: str, findings_report: str, details_excel: str, sheet_name: str = "External") -> None:
    vulnerabilities = severity_counter(technical_report)
    lookup = load_findings_lookup(details_excel, sheet_name=sheet_name)

    doc = Document(findings_report)
    section = "FINDINGS DETAILS"

    insert_index = None
    for i, paragraph in enumerate(doc.paragraphs):
        if section.lower() in paragraph.text.lower():
            insert_index = i

    if insert_index is None:
        print(f"The {section} section was not found.")
        return

    insert_position = doc.paragraphs[insert_index + 1]

    for vuln in reversed(vulnerabilities):
        key = normalize_title(vuln.discoveredName)

        if key in lookup:
            detail = lookup[key]

            # Title
            title_para = insert_position.insert_paragraph_before()
            title_run = title_para.add_run(detail["title"])
            title_run.bold = True
            title_run.font.size = Pt(13)

            # Risk Level
            risk_para = insert_position.insert_paragraph_before()
            risk_run = risk_para.add_run(f"Risk Level: {detail['risk']}")
            risk_run.bold = True
            risk_run.font.size = Pt(11)

            # Description
            desc_para = insert_position.insert_paragraph_before()
            desc_para.add_run("Finding Description: ").bold = True
            desc_para.add_run(detail["description"])

            # Recommendation
            rec_para = insert_position.insert_paragraph_before()
            rec_para.add_run("Recommendation: ").bold = True
            rec_para.add_run(detail["recommendation"])

            # spacing
            insert_position.insert_paragraph_before().add_run("")

        else:
            print(f"⚠ No match found for: {vuln.discoveredName}")

    doc.save(findings_report)
    print("✅ Findings Details saved.")